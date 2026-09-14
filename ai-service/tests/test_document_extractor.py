from pathlib import Path

import pytest

from app.services.document_extractor import (
    UnsupportedDocumentError,
    extract_text,
)


def test_extract_txt(tmp_path: Path):
    file_path = tmp_path / "test.txt"
    file_path.write_text(
        "First line.\n\nSecond line.\n",
        encoding="utf-8",
    )

    result = extract_text(str(file_path))

    assert result == "First line.\nSecond line."


def test_extract_markdown(tmp_path: Path):
    file_path = tmp_path / "test.md"
    file_path.write_text(
        "# Refund Policy\n\nRefunds are allowed within 30 days.",
        encoding="utf-8",
    )

    result = extract_text(str(file_path))

    assert "# Refund Policy" in result
    assert "Refunds are allowed within 30 days." in result


def test_extract_docx(tmp_path: Path):
    from docx import Document

    file_path = tmp_path / "test.docx"

    document = Document()
    document.add_paragraph("Refund Policy")
    document.add_paragraph("Refunds are allowed within 30 days.")
    document.save(file_path)

    result = extract_text(str(file_path))

    assert "Refund Policy" in result
    assert "Refunds are allowed within 30 days." in result


def test_unsupported_extension(tmp_path: Path):
    file_path = tmp_path / "test.csv"
    file_path.write_text("a,b,c", encoding="utf-8")

    with pytest.raises(UnsupportedDocumentError):
        extract_text(str(file_path))